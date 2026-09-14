#!/usr/bin/env python3
"""Check actual embedded image resolution against its displayed PowerPoint size."""

import argparse
import json
from pathlib import Path
from pptx import Presentation
from PIL import Image
from io import BytesIO

EMU_PER_INCH = 914400


def inspect(path, minimum=150):
    prs = Presentation(path)
    rows = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.shape_type == 13:
                continue
            try:
                im = Image.open(BytesIO(sh.image.blob))
                xdpi = im.width / (sh.width / EMU_PER_INCH)
                ydpi = im.height / (sh.height / EMU_PER_INCH)
                ppi = min(xdpi, ydpi)
                rows.append(
                    {
                        "slide": si,
                        "shape": sh.name,
                        "pixels": [im.width, im.height],
                        "effective_ppi": round(ppi, 1),
                        "passes_default": ppi >= minimum,
                    }
                )
            except Exception as e:
                rows.append({"slide": si, "shape": sh.name, "error": str(e), "passes_default": False})
    return {"minimum_ppi": minimum, "images": rows, "passed": all(x.get("passes_default", False) for x in rows)}


def main():
    p = argparse.ArgumentParser()
    p.add_argument("pptx", type=Path)
    p.add_argument("--minimum-ppi", type=int, default=150)
    a = p.parse_args()
    r = inspect(a.pptx, a.minimum_ppi)
    print(json.dumps(r, ensure_ascii=False, indent=2))
    raise SystemExit(0 if r["passed"] else 1)


if __name__ == "__main__":
    main()
