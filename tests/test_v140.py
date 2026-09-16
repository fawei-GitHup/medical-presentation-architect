import base64
import hashlib
import io
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import capture_run  # noqa: E402
import media_index  # noqa: E402
import mpa  # noqa: E402
import perceptual_preflight  # noqa: E402
import source_map_check  # noqa: E402
import workflow_state  # noqa: E402


class V140Tests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="mpa-v140-")
        self.tmp = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def _png(self, path: Path, size=(1800, 1200), color="#247A7A") -> bytes:
        Image.new("RGB", size, color).save(path)
        return path.read_bytes()

    def _deck_with_duplicate_media(self) -> Path:
        image = self.tmp / "large.png"
        self._png(image)
        prs = Presentation()
        for _ in range(2):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(image), Inches(1), Inches(1), width=Inches(5))
        deck = self.tmp / "media.pptx"
        prs.save(deck)
        return deck

    def test_media_index_is_deduplicated_and_low_context(self):
        output = self.tmp / "index"
        report = media_index.index_pptx(self._deck_with_duplicate_media(), output, batch_size=4)
        self.assertEqual(len(report["media"]), 1)
        self.assertEqual({row["slide_id"] for row in report["media"][0]["occurrences"]}, {1, 2})
        serialized = (output / "media_index.json").read_text(encoding="utf-8")
        self.assertNotIn("base64", serialized.lower())
        self.assertNotIn(base64.b64encode((self.tmp / "large.png").read_bytes())[:200].decode(), serialized)
        with Image.open(output / report["media"][0]["thumbnail"]) as thumbnail:
            self.assertLessEqual(max(thumbnail.size), 640)
        self.assertEqual(len(report["contact_sheets"]), 1)

    def test_media_extract_exports_only_selected_bytes_with_matching_hash(self):
        output = self.tmp / "index"
        report = media_index.index_pptx(self._deck_with_duplicate_media(), output)
        extracted = media_index.extract_selected(report, self.tmp / "selected", [report["media"][0]["id"]])
        self.assertEqual(len(extracted), 1)
        path = Path(extracted[0]["path"])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), report["media"][0]["sha256"])

    def test_media_index_rejects_oversized_archive_member_before_decompression(self):
        deck = self.tmp / "oversized.pptx"
        with ZipFile(deck, "w", ZIP_DEFLATED) as archive:
            archive.writestr("ppt/media/oversized.bin", b"x" * 1024)
        with patch.object(media_index, "MAX_MEDIA_BYTES", 128), self.assertRaisesRegex(ValueError, "exceeds"):
            media_index.index_pptx(deck, self.tmp / "oversized-index")

    def test_source_map_rejects_internal_id_without_mapping(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Sources: [SRC1]"
        deck = self.tmp / "bad-citations.pptx"
        prs.save(deck)
        report = source_map_check.check_presentation(deck, [{"id": "SRC1", "authors": "Sannino G", "year": 2015}])
        codes = {issue["code"] for issue in report["issues"]}
        self.assertIn("internal_id_visible", codes)
        self.assertIn("missing_reference_mapping", codes)
        self.assertFalse(report["passed"])

    def test_source_map_accepts_readable_citation_and_complete_reference(self):
        source = {"id": "SRC1", "authors": "Sannino G, Germano F, Arcuri L", "year": 2015, "title": "CEREC review"}
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = "Sannino et al., 2015"
        slide.notes_slide.notes_text_frame.text = "Evidence ledger: [SRC1]"
        references = prs.slides.add_slide(prs.slide_layouts[6])
        references.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1)).text = "参考文献"
        references.shapes.add_textbox(Inches(1), Inches(1.5), Inches(10), Inches(1)).text = (
            "Sannino et al., 2015. CEREC review."
        )
        deck = self.tmp / "good-citations.pptx"
        prs.save(deck)
        report = source_map_check.check_presentation(deck, [source])
        self.assertTrue(report["passed"], report["issues"])
        self.assertNotIn("SRC1", source_map_check.readable_citations(["SRC1"], {"SRC1": source}))

    def test_source_map_reads_grouped_text_and_reference_tables(self):
        source = {"id": "SRC1", "authors": "Sannino G", "year": 2015, "title": "CEREC review"}
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "内部证据 [SRC1]"
        references = prs.slides.add_slide(prs.slide_layouts[6])
        table = references.shapes.add_table(2, 1, Inches(1), Inches(1), Inches(10), Inches(2)).table
        table.cell(0, 0).text = "参考文献"
        table.cell(1, 0).text = "SRC1 — Sannino, 2015. CEREC review."
        deck = self.tmp / "nested-citations.pptx"
        prs.save(deck)
        report = source_map_check.check_presentation(deck, [source])
        codes = {issue["code"] for issue in report["issues"]}
        self.assertIn("internal_id_visible", codes)
        self.assertNotIn("missing_reference_mapping", codes)
        self.assertEqual(report["reference_slides"], [2])

    def test_perceptual_preflight_finds_overlap_and_occluded_connector(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        first = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(2.1), Inches(0.6))
        first.text = "烧结上釉"
        second = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.08), Inches(2), Inches(2.1), Inches(0.6))
        second.text = "当日戴牙"
        third = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.5), Inches(2), Inches(2.1), Inches(0.6))
        third.text = "复核"
        slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.05), Inches(2.3), Inches(4.13), Inches(2.3))
        deck = self.tmp / "overlap.pptx"
        prs.save(deck)
        report = perceptual_preflight.analyze_presentation(deck)
        codes = {issue["code"] for issue in report["issues"]}
        self.assertIn("flow_node_overlap", codes)
        self.assertIn("flow_connector_occluded", codes)
        self.assertIn("flow_connector_missing", codes)

    def test_perceptual_preflight_finds_orphan_empty_card_and_dense_reference(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        orphan = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(1))
        orphan.text = "说明\n字"
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(7), Inches(2))
        card.text = "短"
        references = prs.slides.add_slide(prs.slide_layouts[6])
        title = references.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(4), Inches(0.5))
        title.text = "参考文献"
        box = references.shapes.add_textbox(Inches(0.6), Inches(1.1), Inches(11.5), Inches(5.5))
        box.text = "\n".join(f"Reference {index}: " + "dense citation text " * 6 for index in range(18))
        for paragraph in box.text_frame.paragraphs:
            paragraph.font.size = Pt(8)
        deck = self.tmp / "perceptual.pptx"
        prs.save(deck)
        report = perceptual_preflight.analyze_presentation(deck)
        codes = {issue["code"] for issue in report["issues"]}
        self.assertIn("single_character_line", codes)
        self.assertIn("large_low_information_card", codes)
        self.assertIn("reference_font_too_small", codes)
        self.assertIn("dense_slide", codes)

    def test_stage_resume_invalidates_changed_artifact_and_dependents(self):
        project = self.tmp / "project"
        project.mkdir()
        brief = project / "brief.json"
        brief.write_text("v1", encoding="utf-8")
        research = project / "research.md"
        research.write_text("evidence", encoding="utf-8")
        workflow_state.record_stage(project, "intake", "completed", [Path("brief.json")])
        workflow_state.record_stage(project, "research", "completed", [Path("research.md")])
        self.assertIn("narrative", workflow_state.resume_plan(project)["ready"])
        brief.write_text("v2", encoding="utf-8")
        report = workflow_state.resume_plan(project)
        self.assertIn("intake", report["stale"])
        self.assertIn("research", report["stale"])

    def test_parallel_stage_updates_do_not_overwrite_each_other(self):
        project = self.tmp / "parallel-project"
        project.mkdir()
        phases = ["design_audit", "media", "assets", "research", "doctor", "design_system"]
        code = (
            "import sys;from pathlib import Path;"
            "sys.path.insert(0,sys.argv[1]);import workflow_state;"
            "workflow_state.record_stage(Path(sys.argv[2]),sys.argv[3],'completed',[Path(sys.argv[4])])"
        )
        processes = []
        for phase in phases:
            artifact = project / f"{phase}.txt"
            artifact.write_text(phase, encoding="utf-8")
            processes.append(
                subprocess.Popen(
                    [sys.executable, "-c", code, str(ROOT / "scripts"), str(project), phase, artifact.name],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                )
            )
        for process in processes:
            stdout, stderr = process.communicate(timeout=30)
            self.assertEqual(process.returncode, 0, stdout + stderr)
        state = workflow_state._load(project)
        self.assertEqual(set(state["stages"]), set(phases))

    def test_capture_scrub_removes_data_uri_and_preserves_audit_metadata(self):
        buffer = io.BytesIO()
        Image.new("RGB", (11, 7), "red").save(buffer, format="PNG")
        encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
        cleaned, records = capture_run.scrub_text("payload=data:image/png;base64," + encoded)
        self.assertNotIn(encoded, cleaned)
        self.assertIn("dimensions=11x7", cleaned)
        self.assertIn("sha256=", cleaned)
        self.assertEqual(records[0]["kind"], "data_uri")

    def test_capture_log_redacts_credentials_from_audited_command(self):
        redacted = capture_run.redact_command(["tool", "--api-key", "secret-1", "--token=secret-2", "safe"])
        self.assertEqual(redacted, ["tool", "--api-key", "[REDACTED]", "--token=[REDACTED]", "safe"])

    def test_generated_deck_uses_readable_citations_and_visible_flow_gaps(self):
        project = self.tmp / "demo"
        shutil.copytree(ROOT / "examples/demo", project)
        deck = mpa.build_pptx(project)
        prs = Presentation(str(deck))
        visible_text = "\n".join(
            getattr(shape, "text", "") for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)
        )
        self.assertNotRegex(visible_text, r"\bSRC\d+\b")
        report = perceptual_preflight.analyze_presentation(deck)
        codes = {issue["code"] for issue in report["issues"]}
        self.assertNotIn("flow_node_overlap", codes)
        self.assertNotIn("flow_connector_occluded", codes)


if __name__ == "__main__":
    unittest.main()
