import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import mpa  # noqa: E402
from pptx_lint import lint  # noqa: E402


class MpaTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="mpa-tests-")
        self.tmp = Path(self.temp.name)
        self.project = self.tmp / "project"
        shutil.copytree(
            ROOT / "examples/demo",
            self.project,
            ignore=shutil.ignore_patterns("build", "render", "final", "history", "*.pyc", "__pycache__"),
        )

    def tearDown(self):
        self.temp.cleanup()

    def test_01_fuzzy_route_does_not_invent_nurse_audience(self):
        p = self.tmp / "fuzzy"
        brief = mpa.init_project(p, "fuzzy", "口腔科，有设备也有护士内容", None)
        self.assertEqual(brief["fields"]["audience"]["status"], "unknown")
        notice = mpa.build_notice(brief)
        text = mpa.render_notice(notice)
        self.assertIn("受众", text)
        self.assertNotIn("护士培训需求", text)
        self.assertTrue(any(x["id"] == "audience" for x in notice["questions"]))

    def test_02_known_fields_are_not_reasked(self):
        b = mpa.read_json(self.project / "intake/design_brief.json")
        self.assertFalse(any(q["id"] == "audience" for q in mpa.missing_questions(b)))
        notice = mpa.build_notice(b)
        self.assertTrue(any(x["field"] == "audience" for x in notice["reused_facts"]))

    def test_03_focused_copy_edit_skips_unneeded_full_intake(self):
        b = mpa.init_brief("focused_edit", "Fix punctuation on page 3")
        b["fields"]["topic"].update(value="Edit page 3 punctuation", status="provided")
        b["fields"]["scope"].update(value="Page 3 punctuation only", status="provided")
        questions = mpa.missing_questions(b)
        self.assertFalse(
            any(
                x["id"] in {"audience", "hospital", "template_brand", "network_policy", "privacy_constraints"}
                for x in questions
            )
        )

    def test_focused_edit_gate_does_not_require_full_training_intake(self):
        project = self.tmp / "focused-project"
        brief = mpa.init_project(project, "focused_edit", "Change page 3 punctuation", None)
        brief["fields"]["scope"].update(value="仅调整第 3 页标点", status="provided")
        mpa.write_json(project / "intake/design_brief.json", brief)
        mpa.update_brief_hash(project)
        errs = mpa.cross_errors(project)
        self.assertFalse(
            any("intake incomplete: audience" in e or "intake incomplete: network_policy" in e for e in errs)
        )

    def test_04_safe_defaults_are_disclosed_and_nonblocking(self):
        b = mpa.read_json(self.project / "intake/design_brief.json")
        self.assertEqual(b["fields"]["slide_count_policy"]["status"], "defaulted")
        new_brief = mpa.init_brief("fuzzy", "A brief topic")
        self.assertIn("中文为主", new_brief["fields"]["language"]["value"])
        self.assertEqual(new_brief["fields"]["language"]["status"], "defaulted")
        self.assertTrue(any(x["field"] == "slide_count_policy" for x in mpa.build_notice(b)["assumptions"]))

    def test_05_old_parameters_are_observations_not_claims(self):
        inventory = mpa.read_json(self.project / "intake/content_inventory.json")
        inventory["slides"] = [
            {
                "slide_id": "S017",
                "page": 17,
                "tags": ["materials"],
                "visible_summary": "Unverified source-deck value: 700 MPa",
                "sources": [],
                "notes_status": "missing",
                "observations": ["Observed in supplied slide; not verified"],
            }
        ]
        mpa.write_json(self.project / "intake/content_inventory.json", inventory)
        self.assertEqual(
            mpa.read_json(self.project / "claims.json"),
            [x for x in mpa.read_json(self.project / "claims.json") if x["kind"] == "synthetic"],
        )
        self.assertEqual(mpa.schema_errors(self.project), [])

    def test_06_mixed_audience_is_preserved_in_brief(self):
        b = mpa.read_json(self.project / "intake/design_brief.json")
        b["fields"]["audience"]["value"] = "Doctors and nurses attending a shared workflow workshop"
        self.assertIn("and", b["fields"]["audience"]["value"])
        self.assertEqual(b["fields"]["audience"]["status"], "provided")

    def test_07_discussion_only_build_is_blocked(self):
        b = mpa.read_json(self.project / "intake/design_brief.json")
        b["status"] = "discussion_only"
        mpa.write_json(self.project / "intake/design_brief.json", b)
        with self.assertRaisesRegex(RuntimeError, "discussion only"):
            brief = mpa.read_json(self.project / "intake/design_brief.json")
            if brief.get("status") == "discussion_only":
                raise RuntimeError("This brief is for discussion only; build is not authorized.")

    def test_08_no_fixed_title_formula(self):
        skill = (ROOT / "rules/anti-ai-slop.md").read_text(encoding="utf-8")
        self.assertIn("标题应直指主题或有证据的结论", skill)
        self.assertNotIn("每个标题必须是结论句", skill)

    def test_09_blockers_name_affected_stages(self):
        b = mpa.init_brief("fuzzy", "口腔科")
        qs = mpa.missing_questions(b)
        audience = next(q for q in qs if q["id"] == "audience")
        self.assertEqual(audience["blocking_scope"], ["narrative", "planning"])

    def test_10_notice_and_prompt_share_brief_identity(self):
        h = mpa.compile_prompt(self.project)
        brief = mpa.read_json(self.project / "intake/design_brief.json")
        notice = mpa.read_json(self.project / "intake/user_notice.json")
        prompt = (self.project / "intake/execution_prompt.md").read_text(encoding="utf-8")
        self.assertEqual(h, brief["brief_hash"])
        self.assertEqual(notice["brief_hash"], h)
        self.assertEqual(notice["brief_version"], brief["version"])
        self.assertIn(h, prompt)

    def test_11_notice_prepared_is_not_sent_or_approved(self):
        mpa.compile_prompt(self.project)
        notice = mpa.read_json(self.project / "intake/user_notice.json")
        self.assertEqual(notice["delivery_status"], "prepared")
        self.assertIsNone(notice["message_id"])
        self.assertFalse(notice["requires_response"] is False and notice["delivery_status"] == "sent")

    def test_12_notice_lists_reused_audience(self):
        mpa.compile_prompt(self.project)
        notice = mpa.read_json(self.project / "intake/user_notice.json")
        self.assertIn("audience", [x["field"] for x in notice["reused_facts"]])

    def test_13_unknown_audience_remains_unknown_in_explanation(self):
        b = mpa.init_brief("fuzzy", "口腔科")
        self.assertEqual(b["fields"]["audience"]["status"], "unknown")
        self.assertIn("not yet confirmed", mpa.render_notice(mpa.build_notice(b)))

    def test_14_notice_names_narrow_edit_scope(self):
        b = mpa.init_brief("focused_edit", "Only edit page 3 title")
        b["fields"]["scope"].update(value="Page 3 title only", status="provided")
        self.assertIn("Page 3 title only", mpa.render_notice(mpa.build_notice(b)))

    def test_15_defaults_are_marked_as_temporary(self):
        txt = mpa.render_notice(mpa.build_notice(mpa.read_json(self.project / "intake/design_brief.json")))
        self.assertIn("暂定建议", txt)

    def test_16_clear_ready_notice_does_not_add_confirmation_round(self):
        n = mpa.build_notice(mpa.read_json(self.project / "intake/design_brief.json"))
        self.assertFalse(n["requires_response"])
        self.assertIn("Continue", n["next_action"])

    def test_17_medical_unknown_is_localized_to_clinical_claims(self):
        b = mpa.init_brief("fuzzy", "Dental bonding slide")
        q = next(x for x in mpa.missing_questions(b) if x["id"] == "network_policy")
        self.assertIn("research", q["blocking_scope"])

    def test_18_user_answer_versions_brief_notice_prompt(self):
        cmd = [
            sys.executable,
            str(ROOT / "scripts/mpa.py"),
            "set-field",
            str(self.project),
            "--set",
            'audience="Mixed clinicians"',
            "--origin",
            "user_message",
            "--locator",
            "current answer",
        ]
        subprocess.run(cmd, check=True, capture_output=True, text=True)
        b = mpa.read_json(self.project / "intake/design_brief.json")
        n = mpa.read_json(self.project / "intake/user_notice.json")
        prompt = (self.project / "intake/execution_prompt.md").read_text(encoding="utf-8")
        self.assertEqual(b["version"], 2)
        self.assertEqual(n["brief_version"], 2)
        self.assertEqual(n["brief_hash"], b["brief_hash"])
        self.assertIn(b["brief_hash"], prompt)
        self.assertEqual(len(b["history"]), 1)

    def test_19_generated_notice_never_claims_to_have_been_sent(self):
        n = mpa.build_notice(mpa.read_json(self.project / "intake/design_brief.json"))
        self.assertEqual(n["delivery_status"], "prepared")

    def test_20_focused_notice_stays_short(self):
        b = mpa.init_brief("focused_edit", "Change the page 3 heading")
        text = mpa.render_notice(mpa.build_notice(b))
        self.assertLess(len(text), 1100)

    def test_notice_delivery_record_does_not_stale_content_review(self):
        mpa.compile_prompt(self.project)
        before = mpa.project_fingerprint(self.project)
        mpa.mark_notice_sent(self.project, "codex_conversation", "intake")
        after = mpa.project_fingerprint(self.project)
        self.assertEqual(before, after)
        notice = mpa.read_json(self.project / "intake/user_notice.json")
        self.assertEqual(notice["delivery_status"], "sent")
        self.assertEqual(notice["delivery_channel"], "codex_conversation")

    def test_export_requires_user_notice_to_be_shown(self):
        mpa.compile_prompt(self.project)
        with self.assertRaisesRegex(RuntimeError, "Export blocked"):
            mpa.export_project(self.project)
        report = mpa.read_json(self.project / "qa_report.json")
        self.assertEqual(report["status"], "needs_user_notice")
        self.assertTrue(any("final delivery notice has not been prepared" in e for e in report["errors"]))

    def test_delivery_notice_is_bound_to_current_brief(self):
        notice = mpa.prepare_delivery_notice(self.project)
        self.assertEqual(notice["notice_kind"], "delivery")
        self.assertEqual(notice["delivery_status"], "prepared")
        self.assertEqual(notice["questions"], [])
        mpa.mark_notice_sent(self.project, "codex_conversation", "delivery")
        self.assertEqual(mpa.cross_errors(self.project), [])

    def test_sent_notice_requires_delivery_channel(self):
        mpa.compile_prompt(self.project)
        notice = mpa.read_json(self.project / "intake/user_notice.json")
        notice.update(delivery_status="sent", delivery_channel=None)
        self.assertTrue(mpa.schema_validate("user-notice", notice))

    def test_schemas_all_load(self):
        from jsonschema import Draft202012Validator

        for p in (ROOT / "schemas").glob("*.schema.json"):
            Draft202012Validator.check_schema(mpa.read_json(p))

    def test_synthetic_claims_are_not_recorded_as_human_verified(self):
        claims = mpa.read_json(self.project / "claims.json")
        self.assertTrue(all(c["kind"] == "synthetic" and c["status"] == "synthetic" for c in claims))
        self.assertTrue(all("reviewer" not in c for c in claims))

    def test_demo_builds_editable_shapes_and_page_notes(self):
        out = mpa.build_pptx(self.project)
        prs = Presentation(str(out))
        self.assertEqual(len(prs.slides), 6)
        self.assertTrue(all(s.notes_slide.notes_text_frame.text.strip() for s in prs.slides))
        self.assertTrue(any(sh.has_chart for s in prs.slides for sh in s.shapes))
        self.assertTrue(any(sh.has_table for s in prs.slides for sh in s.shapes))
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        self.assertTrue(any(sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE for s in prs.slides for sh in s.shapes))
        nodes = [sh for sh in prs.slides[1].shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
        self.assertEqual(len(nodes), 3)
        self.assertTrue(all(sh.width > 0 and sh.height > 0 for sh in nodes))
        self.assertTrue(lint(out)["passed"])

    def test_unverified_numeric_claim_is_blocked(self):
        claims = mpa.read_json(self.project / "claims.json")
        claims[0].update(kind="numeric", status="pending", text="A clinical effect of 30%")
        mpa.write_json(self.project / "claims.json", claims)
        # Cross-field verifier catches the evidence issue without depending on a successful render/review.
        self.assertIn("unverified claim: DEMO_1", mpa.machine_report(self.project)["errors"])

    def test_image_overlay_detector_finds_hidden_words(self):
        image = self.tmp / "cover.png"
        Image.new("RGB", (400, 200), "white").save(image)
        p = Presentation()
        s = p.slides.add_slide(p.slide_layouts[6])
        t = s.shapes.add_textbox(100000, 100000, 3000000, 1000000)
        t.text = "covered clinical instructions"
        s.shapes.add_picture(str(image), 100000, 100000, width=3000000, height=1000000)
        deck = self.tmp / "covered.pptx"
        p.save(deck)
        report = lint(deck)
        self.assertTrue(any("covered by" in x for x in report["errors"]))

    def test_hidden_page_detector_reads_ooxml(self):
        p = Presentation()
        p.slides.add_slide(p.slide_layouts[6])
        src = self.tmp / "plain.pptx"
        p.save(src)
        out = self.tmp / "hidden.pptx"
        with ZipFile(src) as zin, ZipFile(out, "w", ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "ppt/slides/slide1.xml":
                    root = ET.fromstring(data)
                    root.set("show", "0")
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                zout.writestr(item, data)
        self.assertTrue(any("hidden" in e for e in lint(out)["errors"]))

    def test_final_status_requires_real_manual_review_and_render(self):
        report = mpa.machine_report(self.project)
        self.assertFalse(report["passed"])
        self.assertIn("manual review missing", report["errors"])
        self.assertIn("render manifest missing; render every slide before review", report["errors"])

    def test_potential_user_data_is_not_in_zip(self):
        import secrets

        private_root = ROOT / "private"
        private_root.mkdir(exist_ok=True)
        token = secrets.token_hex(16)
        private_dir = private_root / token
        private_dir.mkdir()
        private = private_dir / "case-record.txt"
        sentinel = ("case-" + token).encode("ascii")
        private.write_bytes(sentinel)
        zipout = self.tmp / "safe.zip"
        try:
            subprocess.run(
                [sys.executable, str(ROOT / "scripts/package.py"), "--output", str(zipout)],
                check=True,
                capture_output=True,
                text=True,
            )
            with ZipFile(zipout) as z:
                names = z.namelist()
                data = b" ".join(z.read(n) for n in names if not n.endswith("/"))
                self.assertFalse(any("/build/" in n or "/render/" in n or "/private/" in n for n in names))
                self.assertNotIn(sentinel, data)
        finally:
            private.unlink(missing_ok=True)
            private_dir.rmdir()
            try:
                private_root.rmdir()
            except OSError:
                pass

    def test_installer_refuses_collision_and_uninstall_preserves_edits(self):
        parent = self.tmp / "skills"
        cmd = [
            sys.executable,
            str(ROOT / "scripts/install.py"),
            "install",
            "--agent",
            "generic",
            "--target",
            str(parent),
        ]
        subprocess.run(cmd, check=True, capture_output=True, text=True)
        dest = parent / "medical-presentation-architect"
        edited = dest / "SKILL.md"
        edited.write_text(edited.read_text(encoding="utf-8") + "\nuser local note\n", encoding="utf-8")
        with self.assertRaises(subprocess.CalledProcessError):
            subprocess.run(cmd, check=True, capture_output=True, text=True)
        subprocess.run(
            [
                sys.executable,
                str(ROOT / "scripts/install.py"),
                "uninstall",
                "--agent",
                "generic",
                "--target",
                str(parent),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        self.assertIn("user local note", edited.read_text(encoding="utf-8"))
        self.assertTrue(dest.exists())


if __name__ == "__main__":
    unittest.main()
