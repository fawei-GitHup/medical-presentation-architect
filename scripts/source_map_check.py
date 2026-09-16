#!/usr/bin/env python3
"""Check audience-facing citations against a deck's evidence ledger.

Internal source IDs remain useful in speaker notes and machine ledgers, but they
must not be the only citation an audience sees.  This module deliberately does
not claim that a source supports a medical statement; it only checks display
and referential integrity.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path


SOURCE_ID_RE = re.compile(r"\bSRC\s*[-_]?\s*(\d+)\b", re.I)
REFERENCE_HINT_RE = re.compile(r"参考文献|参考资料|主要来源|references?|bibliography", re.I)


def _natural_id(value: str) -> tuple[int, str]:
    match = re.search(r"(\d+)$", value)
    return (int(match.group(1)) if match else 10**9, value)


def _author_parts(authors) -> list[str]:
    if isinstance(authors, list):
        raw = [str(value).strip() for value in authors]
    else:
        text = str(authors or "").strip()
        if not text:
            return []
        raw = [part.strip() for part in re.split(r"\s*[,;；]\s*", text) if part.strip()]
    return [part for part in raw if not re.match(r"^(et\s+al\.?|等)$", part, re.I)]


def _family_name(author: str) -> str:
    author = re.sub(r"\s+", " ", author.strip())
    if not author:
        return ""
    if re.search(r"\b(organization|association|society|university|institute|ministry|commission)\b", author, re.I):
        return author
    return author.split()[0]


def short_citation(source: dict) -> str:
    """Return a compact human-readable citation without leaking an internal ID."""
    authors = _author_parts(source.get("authors") or source.get("author") or source.get("institution"))
    year = str(source.get("year") or "n.d.")
    if len(authors) >= 3:
        label = f"{_family_name(authors[0])} et al."
    elif len(authors) == 2:
        label = f"{_family_name(authors[0])} & {_family_name(authors[1])}"
    elif authors:
        label = _family_name(authors[0])
    else:
        title = str(source.get("title") or "Untitled source").strip()
        label = title if len(title) <= 34 else title[:31].rstrip() + "…"
    return f"{label}, {year}"


def readable_citations(source_ids: list[str], sources: dict[str, dict], max_items: int = 6) -> str:
    ordered = sorted(dict.fromkeys(source_ids), key=_natural_id)
    labels = [short_citation(sources[sid]) for sid in ordered if sid in sources]
    if len(labels) > max_items:
        labels = labels[:max_items] + [f"另 {len(labels) - max_items} 项（见参考文献）"]
    return "; ".join(labels)


def _shape_text(shape) -> str:
    chunks = []
    if getattr(shape, "has_text_frame", False):
        chunks.append(str(shape.text or ""))
    if getattr(shape, "has_table", False):
        chunks.extend(cell.text for row in shape.table.rows for cell in row.cells)
    children = getattr(shape, "shapes", None)
    if children is not None:
        chunks.extend(_shape_text(child) for child in children)
    return "\n".join(chunk for chunk in chunks if chunk)


def _slide_text(slide) -> str:
    """Collect visible text from text frames, tables, and nested groups."""
    return "\n".join(filter(None, (_shape_text(shape) for shape in slide.shapes)))


def _note_text(slide) -> str:
    try:
        return slide.notes_slide.notes_text_frame.text or ""
    except (AttributeError, KeyError, ValueError):
        return ""


def _source_tokens(source: dict) -> list[str]:
    tokens = [short_citation(source)]
    title = str(source.get("title") or "").strip()
    if title:
        tokens.append(title[:48])
    for field in ("doi", "pmid"):
        if source.get(field):
            tokens.append(str(source[field]))
    return [token for token in tokens if len(token) >= 3]


def check_presentation(pptx_path: Path, sources: list[dict] | None = None) -> dict:
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    source_map = {str(source.get("id")): source for source in (sources or []) if source.get("id")}
    slide_records = []
    used_ids: list[str] = []
    visible_ids: list[tuple[int, str]] = []
    reference_slides: list[int] = []
    reference_text = []
    for number, slide in enumerate(prs.slides, 1):
        body = _slide_text(slide)
        notes = _note_text(slide)
        body_ids = [f"SRC{match}" for match in SOURCE_ID_RE.findall(body)]
        note_ids = [f"SRC{match}" for match in SOURCE_ID_RE.findall(notes)]
        is_reference = bool(REFERENCE_HINT_RE.search(body[:500]))
        if is_reference:
            reference_slides.append(number)
            reference_text.append(body)
        else:
            visible_ids.extend((number, sid) for sid in body_ids)
        used_ids.extend(body_ids + note_ids)
        slide_records.append(
            {
                "slide_id": number,
                "visible_source_ids": body_ids,
                "notes_source_ids": note_ids,
                "reference_slide": is_reference,
            }
        )

    joined_references = "\n".join(reference_text)
    explicit_reference_ids = [f"SRC{match}" for match in SOURCE_ID_RE.findall(joined_references)]
    explicit_counts = Counter(explicit_reference_ids)
    issues = []
    for slide_id, sid in visible_ids:
        issues.append(
            {
                "code": "internal_id_visible",
                "severity": "error",
                "slide_id": slide_id,
                "shape": None,
                "evidence": f"{sid} is visible to the audience",
                "suggestion": "Replace the internal ID with an author-year short citation; retain the ID in notes.",
            }
        )
    ordered_used = sorted(dict.fromkeys(used_ids), key=_natural_id)
    for sid in ordered_used:
        if source_map and sid not in source_map:
            issues.append(
                {
                    "code": "unknown_source_id",
                    "severity": "error",
                    "slide_id": None,
                    "shape": None,
                    "evidence": f"{sid} is used but absent from sources.json",
                    "suggestion": "Correct the ID or add the missing source ledger entry.",
                }
            )
            continue
        mapped = sid in explicit_reference_ids
        if not mapped and sid in source_map:
            folded = joined_references.casefold()
            mapped = any(token.casefold() in folded for token in _source_tokens(source_map[sid]))
        if not mapped:
            issues.append(
                {
                    "code": "missing_reference_mapping",
                    "severity": "error",
                    "slide_id": None,
                    "shape": None,
                    "evidence": f"{sid} is used but has no complete mapping on a reference/appendix slide",
                    "suggestion": "Add one complete reference entry matching the author-year citation or an explicit internal-ID mapping.",
                }
            )
    for sid, count in sorted(explicit_counts.items(), key=lambda item: _natural_id(item[0])):
        if count > 1:
            issues.append(
                {
                    "code": "duplicate_reference_mapping",
                    "severity": "warning",
                    "slide_id": None,
                    "shape": None,
                    "evidence": f"{sid} appears {count} times in the reference mapping",
                    "suggestion": "Keep one canonical reference entry for each source ID.",
                }
            )
    explicit_numbers = [_natural_id(sid)[0] for sid in explicit_reference_ids]
    if explicit_numbers and explicit_numbers != sorted(explicit_numbers):
        issues.append(
            {
                "code": "unordered_reference_ids",
                "severity": "warning",
                "slide_id": reference_slides[0] if reference_slides else None,
                "shape": None,
                "evidence": f"reference IDs are ordered {explicit_numbers}",
                "suggestion": "Order reference mappings numerically or by first appearance and use that order consistently.",
            }
        )
    return {
        "format": "mpa-source-map-check-v1",
        "pptx": str(pptx_path),
        "passed": not any(issue["severity"] == "error" for issue in issues),
        "reference_slides": reference_slides,
        "used_source_ids": ordered_used,
        "issues": issues,
        "slides": slide_records,
        "semantic_support_reviewed": False,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--sources", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    sources = json.loads(args.sources.read_text(encoding="utf-8")) if args.sources else []
    report = check_presentation(args.pptx, sources)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
