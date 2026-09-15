#!/usr/bin/env python3
"""Extract a compact, editable design fingerprint from a reference PPTX."""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path


def _safe_rgb(color):
    try:
        value = color.rgb
        return f"#{value}" if value else None
    except (AttributeError, TypeError, ValueError):
        return None


def _shape_fill_rgb(shape):
    try:
        return _safe_rgb(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def _composition(slide, index, total, width):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    tables = [shape for shape in slide.shapes if getattr(shape, "has_table", False)]
    charts = [shape for shape in slide.shapes if getattr(shape, "has_chart", False)]
    text = " ".join(shape.text for shape in text_shapes)
    if index == 1:
        return "cover"
    if index == total or any(word in text.lower() for word in ("谢谢", "thank you", "讨论")):
        return "closing"
    if any(word in text.lower() for word in ("参考", "references", "bibliography")):
        return "references"
    if tables:
        return "table"
    if charts:
        return "chart"
    if len(pictures) >= 3:
        return "visual_grid"
    if pictures:
        picture = max(pictures, key=lambda shape: shape.width * shape.height)
        if picture.width >= width * 0.72:
            return "image_hero"
        return "image_text_split" if picture.left < width / 2 else "text_image_split"
    if len(text_shapes) >= 4:
        return "cards_or_steps"
    return "editorial_text"


def audit_pptx_design(source: Path) -> dict:
    from pptx import Presentation

    source = source.expanduser().resolve()
    prs = Presentation(str(source))
    fonts, colors, minimum_fonts = Counter(), Counter(), []
    composition_counts = Counter()
    representative = defaultdict(list)
    geometry_counts = Counter()
    page_observations = []
    for index, slide in enumerate(prs.slides, 1):
        composition = _composition(slide, index, len(prs.slides), prs.slide_width)
        composition_counts[composition] += 1
        if len(representative[composition]) < 3:
            representative[composition].append(index)
        geometry = []
        page_fonts = []
        for shape in slide.shapes:
            geometry.append(
                (
                    int(shape.left / prs.slide_width * 4),
                    int(shape.top / prs.slide_height * 3),
                    int(shape.width / prs.slide_width * 4),
                    int(shape.height / prs.slide_height * 3),
                )
            )
            fill_rgb = _shape_fill_rgb(shape)
            line_rgb = _safe_rgb(getattr(getattr(shape, "line", None), "color", None))
            if fill_rgb:
                colors[fill_rgb.upper()] += 1
            if line_rgb:
                colors[line_rgb.upper()] += 1
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts[run.font.name] += 1
                        if run.font.size:
                            pt = round(run.font.size.pt, 1)
                            minimum_fonts.append(pt)
                            page_fonts.append(pt)
                        font_rgb = _safe_rgb(run.font.color)
                        if font_rgb:
                            colors[font_rgb.upper()] += 1
        geometry_counts[tuple(sorted(geometry))] += 1
        if page_fonts and min(page_fonts) < 12:
            page_observations.append(f"第 {index} 页存在小于 12pt 的文字")
    repeated_geometry_pages = sum(count for count in geometry_counts.values() if count >= 3)
    preserve = []
    if colors:
        preserve.append("保留高频品牌色与中性色之间的对比关系")
    if len(composition_counts) >= 5:
        preserve.append("保留多种页面构图带来的节奏变化")
    if any(name in composition_counts for name in ("image_hero", "visual_grid", "image_text_split", "text_image_split")):
        preserve.append("保留图像主导页和图文分栏页的视觉层级")
    repair = []
    if repeated_geometry_pages:
        repair.append(f"减少重复几何模板；约 {repeated_geometry_pages} 页落入高频重复构图")
    if minimum_fonts and min(minimum_fonts) < 12:
        repair.append("放大过小正文或脚注，并拆分页内过密内容")
    if len(composition_counts) < 5 and len(prs.slides) >= 15:
        repair.append("扩展构图家族，避免整套演示文稿只有少数模板")
    if not repair:
        repair.append("逐页核对截图可读性、内容密度和风险信息层级")
    observations = [
        f"共 {len(prs.slides)} 页；识别到 {len(composition_counts)} 种主要构图",
        f"最常见构图：{composition_counts.most_common(1)[0][0]}（{composition_counts.most_common(1)[0][1]} 页）",
        *page_observations[:8],
    ]
    return {
        "source": str(source),
        "slide_size": {"width": round(prs.slide_width / 914400, 3), "height": round(prs.slide_height / 914400, 3)},
        "fonts": [name for name, _ in fonts.most_common(10)] or ["Arial"],
        "colors": [value for value, _ in colors.most_common(12)] or ["#1D2B36", "#FFFFFF"],
        "representative_slides": dict(representative),
        "composition_counts": dict(composition_counts),
        "preserve": preserve or ["保留现有标题与正文的信息层级"],
        "repair": repair,
        "observations": observations,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = audit_pptx_design(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(args.output)


if __name__ == "__main__":
    main()
