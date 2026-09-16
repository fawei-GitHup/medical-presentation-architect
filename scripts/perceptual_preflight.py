#!/usr/bin/env python3
"""PPTX-level perceptual heuristics for presentation QA.

The checks are intentionally conservative and auditable.  They flag geometry
and typography that needs human review; they do not replace rendered-page
inspection or clinical review.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

EMU_PER_INCH = 914400
REFERENCE_RE = re.compile(r"参考文献|参考资料|references?|bibliography", re.I)


def _inch(value) -> float:
    return float(value) / EMU_PER_INCH


def _box(shape) -> tuple[float, float, float, float]:
    return tuple(_inch(value) for value in (shape.left, shape.top, shape.width, shape.height))


def _shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False):
        return str(shape.text or "").strip()
    if getattr(shape, "has_table", False):
        return "\n".join(cell.text for row in shape.table.rows for cell in row.cells).strip()
    return ""


def _deep_shape_text(shape) -> str:
    chunks = [_shape_text(shape)]
    children = getattr(shape, "shapes", None)
    if children is not None:
        chunks.extend(_deep_shape_text(child) for child in children)
    return "\n".join(chunk for chunk in chunks if chunk)


def _font_sizes(shape) -> list[float]:
    sizes = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            if paragraph.font.size:
                sizes.append(float(paragraph.font.size.pt))
            for run in paragraph.runs:
                if run.font.size:
                    sizes.append(float(run.font.size.pt))
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    if paragraph.font.size:
                        sizes.append(float(paragraph.font.size.pt))
                    for run in paragraph.runs:
                        if run.font.size:
                            sizes.append(float(run.font.size.pt))
    return sizes


def _issue(code, severity, slide_id, shape, evidence, suggestion):
    return {
        "code": code,
        "severity": severity,
        "slide_id": slide_id,
        "shape": shape,
        "evidence": evidence,
        "suggestion": suggestion,
    }


def _intersection(a, b) -> tuple[float, float]:
    return max(0.0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])), max(
        0.0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    )


def _is_connector(shape) -> bool:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        return shape.shape_type == MSO_SHAPE_TYPE.LINE
    except (AttributeError, ValueError):
        return False


def _is_node(shape) -> bool:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            return False
    except (AttributeError, ValueError):
        return False
    x, y, w, h = _box(shape)
    text = _shape_text(shape)
    return bool(text) and y >= 1.1 and 0.25 <= h <= 1.55 and 0.55 <= w <= 4.0 and len(text) <= 80


def _line_endpoints(shape) -> tuple[tuple[float, float], tuple[float, float]]:
    x, y, w, h = _box(shape)
    return (x, y), (x + w, y + h)


def _connector_for_pair(lines, left, right):
    left_center = (left[0] + left[2] / 2, left[1] + left[3] / 2)
    right_center = (right[0] + right[2] / 2, right[1] + right[3] / 2)
    candidates = []
    for line in lines:
        p1, p2 = _line_endpoints(line)
        lx, rx = sorted((p1[0], p2[0]))
        cy = (p1[1] + p2[1]) / 2
        if lx <= left_center[0] + 0.18 and rx >= right_center[0] - 0.18 and abs(cy - left_center[1]) <= 0.4:
            candidates.append((line, p1, p2))
    return candidates[0] if candidates else None


def _layout_signature(shapes, slide_w, slide_h) -> tuple:
    entries = []
    for shape in shapes:
        if _is_connector(shape):
            continue
        text = _shape_text(shape)
        x, y, w, h = _box(shape)
        if not text and not getattr(shape, "shape_type", None):
            continue
        entries.append(
            (
                int(round(x / slide_w * 10)),
                int(round(y / slide_h * 10)),
                int(round(w / slide_w * 10)),
                int(round(h / slide_h * 10)),
                int(getattr(shape, "shape_type", 0)),
            )
        )
    return tuple(sorted(entries))


def analyze_presentation(pptx_path: Path, rendered_dir: Path | None = None) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(pptx_path))
    slide_w, slide_h = _inch(prs.slide_width), _inch(prs.slide_height)
    slide_area = slide_w * slide_h
    issues = []
    signatures = []
    slide_summaries = []
    for slide_id, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        signatures.append(_layout_signature(shapes, slide_w, slide_h))
        full_text = "\n".join(filter(None, (_deep_shape_text(shape) for shape in shapes)))
        reference_slide = bool(REFERENCE_RE.search(full_text[:700]))
        total_chars = sum(len(re.sub(r"\s+", "", _deep_shape_text(shape))) for shape in shapes)
        cards = []
        lines = [shape for shape in shapes if _is_connector(shape)]
        nodes = [shape for shape in shapes if _is_node(shape)]
        for shape in shapes:
            name = str(getattr(shape, "name", "") or f"shape-{shape.shape_id}")
            x, y, w, h = _box(shape)
            text = _shape_text(shape)
            area = max(0.0, w * h)
            if getattr(shape, "shapes", None) is not None:
                issues.append(
                    _issue(
                        "group_shape_requires_render_review",
                        "warning",
                        slide_id,
                        name,
                        "grouped or nested shapes use transform semantics that are not fully represented by top-level geometry checks",
                        "Inspect the rendered slide; ungroup before geometry-critical editing when practical.",
                    )
                )
            if x < -0.01 or y < -0.01 or x + w > slide_w + 0.01 or y + h > slide_h + 0.01:
                issues.append(
                    _issue(
                        "outside_slide_bounds",
                        "error",
                        slide_id,
                        name,
                        f"box=({x:.2f},{y:.2f},{w:.2f},{h:.2f}) on {slide_w:.2f}×{slide_h:.2f} in canvas",
                        "Move or resize the element so its full bounding box remains on the slide.",
                    )
                )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and (w < 0.75 or h < 0.55 or area < 0.55):
                issues.append(
                    _issue(
                        "image_too_small",
                        "warning",
                        slide_id,
                        name,
                        f"image is only {w:.2f}×{h:.2f} in",
                        "Enlarge it enough to read its information or remove it if it is decorative.",
                    )
                )
            if not text:
                continue
            chars = len(re.sub(r"\s+", "", text))
            sizes = _font_sizes(shape)
            min_font = min(sizes) if sizes else None
            hard_lines = [line.strip() for line in text.splitlines() if line.strip()]
            for line in hard_lines:
                clean = re.sub(r"[\s•·☐☑□✓\-–—:：,，。.;；()（）\[\]]", "", line)
                if len(clean) == 1 and re.search(r"[\u3400-\u9fffA-Za-z]", clean):
                    issues.append(
                        _issue(
                            "single_character_line",
                            "warning",
                            slide_id,
                            name,
                            f"explicit line contains one content character: {line!r}",
                            "Rewrap or widen the text box so a single character is not stranded.",
                        )
                    )
                    break
            if min_font and len(hard_lines) == 1 and chars >= 12 and w > 0:
                capacity = max(2, int((w - 0.18) * 72 / (min_font * 0.92)))
                if chars > capacity and chars % capacity == 1:
                    issues.append(
                        _issue(
                            "possible_single_character_wrap",
                            "warning",
                            slide_id,
                            name,
                            f"{chars} compact characters at {min_font:.1f} pt in {w:.2f} in yields an estimated one-character final line",
                            "Inspect the rendered page and adjust copy, width, or font size.",
                        )
                    )
            density = chars / max(area, 0.15)
            if chars >= 90 and (density > 48 or (min_font is not None and min_font < 13)):
                issues.append(
                    _issue(
                        "dense_text_box",
                        "warning",
                        slide_id,
                        name,
                        f"{chars} characters in {area:.2f} in² ({density:.1f}/in²), min font={min_font or 'unknown'} pt",
                        "Split, shorten, or move supporting detail to speaker notes/appendix.",
                    )
                )
            if reference_slide and min_font is not None and min_font < 10.5:
                issues.append(
                    _issue(
                        "reference_font_too_small",
                        "error",
                        slide_id,
                        name,
                        f"reference text uses {min_font:.1f} pt",
                        "Use a readable reference size or distribute references across additional slides.",
                    )
                )
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                cards.append(shape)
                if area / slide_area >= 0.115 and chars <= 20:
                    issues.append(
                        _issue(
                            "large_low_information_card",
                            "warning",
                            slide_id,
                            name,
                            f"card occupies {area / slide_area:.0%} of slide area but contains only {chars} characters",
                            "Reduce the card, add meaningful evidence, or replace the container with direct composition.",
                        )
                    )
        if len(cards) > 7:
            issues.append(
                _issue(
                    "too_many_cards",
                    "warning",
                    slide_id,
                    None,
                    f"{len(cards)} text-bearing auto-shapes on one slide",
                    "Create clearer hierarchy and reduce repeated card containers.",
                )
            )
        if total_chars > (1250 if reference_slide else 720):
            issues.append(
                _issue(
                    "dense_slide",
                    "warning",
                    slide_id,
                    None,
                    f"slide contains {total_chars} compact characters",
                    "Split the slide and keep only audience-critical text in the foreground.",
                )
            )

        # Group likely process nodes into horizontal lanes.
        remaining = sorted(nodes, key=lambda shape: (_box(shape)[1], _box(shape)[0]))
        lanes = []
        for node in remaining:
            box = _box(node)
            center_y = box[1] + box[3] / 2
            lane = next((item for item in lanes if abs(item[0] - center_y) <= 0.34), None)
            if lane:
                lane[1].append(node)
                lane[0] = sum(_box(value)[1] + _box(value)[3] / 2 for value in lane[1]) / len(lane[1])
            else:
                lanes.append([center_y, [node]])
        for _, lane_nodes in lanes:
            lane_nodes.sort(key=lambda shape: _box(shape)[0])
            if len(lane_nodes) < 2:
                continue
            process_like = len(lane_nodes) >= 3 or any(
                abs((_line_endpoints(line)[0][1] + _line_endpoints(line)[1][1]) / 2 - _box(lane_nodes[0])[1] - _box(lane_nodes[0])[3] / 2)
                <= 0.4
                for line in lines
            )
            if not process_like:
                continue
            for left_shape, right_shape in zip(lane_nodes, lane_nodes[1:]):
                left, right = _box(left_shape), _box(right_shape)
                overlap_x, overlap_y = _intersection(left, right)
                gap = right[0] - (left[0] + left[2])
                pair = f"{left_shape.name} → {right_shape.name}"
                if overlap_x > 0.005 and overlap_y > 0.05:
                    issues.append(
                        _issue(
                            "flow_node_overlap",
                            "error",
                            slide_id,
                            pair,
                            f"nodes overlap horizontally by {overlap_x:.3f} in",
                            "Distribute nodes from the available lane width using an explicit minimum gap.",
                        )
                    )
                elif gap < 0.08:
                    issues.append(
                        _issue(
                            "flow_node_gap_too_small",
                            "error",
                            slide_id,
                            pair,
                            f"visible gap is {gap:.3f} in",
                            "Use at least 0.08 in of visible separation, preferably more for a readable connector.",
                        )
                    )
                connection = _connector_for_pair(lines, left, right)
                if not connection:
                    issues.append(
                        _issue(
                            "flow_connector_missing",
                            "warning",
                            slide_id,
                            pair,
                            "no connector spans the adjacent process nodes",
                            "Add a connector whose endpoints terminate near the facing node boundaries.",
                        )
                    )
                else:
                    line, p1, p2 = connection
                    start_x, end_x = sorted((p1[0], p2[0]))
                    hidden_left = max(0.0, left[0] + left[2] - start_x)
                    hidden_right = max(0.0, end_x - right[0])
                    if hidden_left > 0.16 or hidden_right > 0.16 or gap <= 0:
                        issues.append(
                            _issue(
                                "flow_connector_occluded",
                                "error",
                                slide_id,
                                line.name,
                                f"connector is hidden under nodes by about {hidden_left:.2f} in and {hidden_right:.2f} in; gap={gap:.3f} in",
                                "Anchor the connector at node boundaries and preserve a visible inter-node gap.",
                            )
                        )
        slide_summaries.append(
            {
                "slide_id": slide_id,
                "shape_count": len(shapes),
                "text_characters": total_chars,
                "card_count": len(cards),
                "reference_slide": reference_slide,
            }
        )

    run_start = 0
    for index in range(1, len(signatures) + 1):
        if index < len(signatures) and signatures[index] == signatures[run_start] and signatures[index]:
            continue
        run_length = index - run_start
        if run_length >= 3:
            issues.append(
                _issue(
                    "repeated_page_structure",
                    "warning",
                    f"{run_start + 1}-{index}",
                    None,
                    f"{run_length} consecutive slides share the same coarse geometry signature",
                    "Vary composition where the narrative role changes; keep repetition only when it aids comparison.",
                )
            )
        run_start = index

    if rendered_dir is not None and rendered_dir.exists():
        rendered = list(rendered_dir.glob("*.png"))
        if rendered and len(rendered) != len(prs.slides):
            issues.append(
                _issue(
                    "render_page_count_mismatch",
                    "error",
                    None,
                    None,
                    f"{len(prs.slides)} PPTX slides but {len(rendered)} rendered PNG pages",
                    "Render every slide before visual review.",
                )
            )
    return {
        "format": "mpa-perceptual-preflight-v1",
        "pptx": str(pptx_path),
        "passed": not any(issue["severity"] == "error" for issue in issues),
        "issues": issues,
        "slides": slide_summaries,
        "limitations": [
            "Line-wrap and information-density findings are heuristics and require rendered-page confirmation.",
            "Nested group geometry is flagged for rendered review rather than treated as exact child coordinates.",
            "This report does not establish medical correctness, privacy compliance, or source support.",
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--rendered-dir", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = analyze_presentation(args.pptx, args.rendered_dir)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
